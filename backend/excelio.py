# -*- coding: utf-8 -*-
"""Excel 导入导出（openpyxl）"""
import io
import json
from datetime import datetime, date
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import db as D

SELECT_TYPES = ("select", "multi_select", "radio", "checkbox")
HEADER_FILL = "2F5B7C"


def _cell_str(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (list, tuple)):
        return "、".join(str(x) for x in v)
    return v


def export_xlsx(table, records, template=False):
    wb = Workbook()
    ws = wb.active
    ws.title = (table["name"] or "数据")[:31]
    fields = sorted(table["fields"], key=lambda f: (f["sort"], f["id"]))

    thin = Side(style="thin", color="D6DEE6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    # 表头
    for c, f in enumerate(fields, 1):
        cell = ws.cell(row=1, column=c, value=f["label"])
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(26, len(f["label"]) * 2.4 + 4))
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    if not template:
        for r_i, r in enumerate(records, 2):
            for c, f in enumerate(fields, 1):
                cell = ws.cell(row=r_i, column=c, value=_cell_str(r["data"].get(f["key"])))
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if f["key"] == "sn":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            if r_i % 2 == 0:
                for c in range(1, len(fields) + 1):
                    if ws.cell(row=r_i, column=c).fill.fgColor.rgb in (None, "00000000"):
                        ws.cell(row=r_i, column=c).fill = PatternFill("solid", fgColor="F5F8FB")

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    fname = f"{table['name']}_{D.now().strftime('%Y%m%d')}.xlsx"
    from flask import Response
    return Response(
        bio.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f"attachment; filename*=UTF-8''{quote(fname)}"})


def import_xlsx(table, f):
    """附加式导入：按列标题匹配字段，逐行校验追加"""
    fields = sorted(table["fields"], key=lambda x: (x["sort"], x["id"]))
    label_map, key_map = {}, {}
    for f_ in fields:
        label_map[f_["label"].strip().lower()] = f_
        key_map[f_["key"].lower()] = f_

    try:
        wb = load_workbook(f, read_only=True, data_only=True)
    except Exception as e:
        return {"added": 0, "failed": [{"row": 0, "reason": f"文件解析失败: {e}"}],
                "ignored_columns": []}
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"added": 0, "failed": [{"row": 0, "reason": "表格为空"}],
                "ignored_columns": []}

    headers = [str(h or "").strip() for h in rows[0]]
    col_field, ignored = {}, []
    for idx, h in enumerate(headers):
        fl = label_map.get(h.lower())
        if fl is None:
            fl = key_map.get(h.lower())
        if fl is not None:
            col_field[idx] = fl
        elif h:
            ignored.append(h)

    from .tables import do_create_record, merge_options
    added, failed = 0, []
    for r_i, row in enumerate(rows[1:], 2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        data = {}
        for idx, fl in col_field.items():
            if idx < len(row):
                v = row[idx]
                v = _cell_str(v)
                if fl["type"] == "number":
                    try:
                        v = int(float(v)) if str(v) not in ("", None) else ""
                    except (TypeError, ValueError):
                        pass
                if fl["type"] == "switch":
                    v = 1 if str(v).lower() in ("1", "true", "是", "y", "yes", "on") else 0
                data[fl["key"]] = "" if v is None else v
        record, err = do_create_record(table, data, "导入")
        if err:
            failed.append({"row": r_i, "reason": err})
        else:
            added += 1
            merge_options(table, data)

    return {"added": added, "failed": failed, "ignored_columns": ignored}
